"""Build the evidence-linked citrus baseline and group-meeting workbook."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT_DIR = Path(__file__).resolve().parent
SERVER_DIR = OUT_DIR.parent
WORKSPACE = SERVER_DIR.parent
RESULTS = SERVER_DIR / "results"
CODE = SERVER_DIR / "code" / "ultralytics-main-new"
LEGACY_SUMMARY = SERVER_DIR / "柑橘实验全量汇总表_20260825.csv"
OUT_XLSX = OUT_DIR / "柑橘实例分割_基线与组会阶段结果_20260915_含非YOLO.xlsx"
OUT_MD = OUT_DIR / "组会汇报速读_基线78到阶段81_20260915.md"

BLUE = "1F4E78"
MID_BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_ORANGE = "FCE4D6"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="B7B7B7")


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return [{str(k).strip(): str(v).strip() for k, v in row.items()} for row in csv.DictReader(stream)]


def best_yolo_row(path: Path) -> dict[str, float | int | str]:
    rows = read_csv(path)
    if not rows:
        raise ValueError(f"Empty results file: {path}")
    row = max(rows, key=lambda item: float(item["metrics/mAP50-95(M)"]))
    return {
        "epoch": int(float(row["epoch"])),
        "box_p": float(row["metrics/precision(B)"]),
        "box_r": float(row["metrics/recall(B)"]),
        "box_map50": float(row["metrics/mAP50(B)"]),
        "box_map": float(row["metrics/mAP50-95(B)"]),
        "mask_p": float(row["metrics/precision(M)"]),
        "mask_r": float(row["metrics/recall(M)"]),
        "mask_map50": float(row["metrics/mAP50(M)"]),
        "mask_map": float(row["metrics/mAP50-95(M)"]),
        "results": str(path.resolve()),
        "prediction": str((path.parent / "val_batch0_pred.jpg").resolve())
        if (path.parent / "val_batch0_pred.jpg").exists()
        else "",
    }


def find_best_epoch(history_path: Path) -> int:
    rows = json.loads(history_path.read_text(encoding="utf-8"))
    eligible = [row for row in rows if isinstance(row.get("val"), dict) and "mask_ap_50_95" in row["val"]]
    return int(max(eligible, key=lambda row: row["val"]["mask_ap_50_95"])["epoch"])


legacy_rows = read_csv(LEGACY_SUMMARY)
legacy_by_run = {row["实验标识"]: row for row in legacy_rows}


def legacy_value(run: str, key: str, default: str = "") -> str:
    return legacy_by_run.get(run, {}).get(key, default)


def optional_float(value: object) -> float | None:
    return float(value) if value is not None and str(value).strip() else None


baseline_specs = [
    ("YOLOv8n-seg", "YOLO 官方轻量基线", "001_1_yolov8-seg_adamw"),
    ("YOLOv9c-seg", "YOLO 历史精度参照", "001_2_yolov9c-seg_adamw"),
    ("YOLO11n-seg", "主消融基线", "001_3_yolo11-seg_adamw"),
    ("YOLO12n-seg", "YOLO 代际参照", "001_4_yolo12-seg_adamw"),
    ("YOLO26n-seg", "YOLO 代际参照", "001_5_yolo26-seg_adamw"),
    ("YOLO11-StarNet-s1", "替换骨干对照", "002_1_yolo11-starnet_s1"),
    ("YOLO11-StarNet-s2", "替换骨干对照", "002_2_yolo11-starnet_s2"),
    ("YOLO11-MobileNetV4", "移动端骨干对照", "003_1_yolo11-mobilenetv4"),
    ("YOLO11-MANO", "注意力对照", "004_1_yolo11-mano(attn)"),
]


baseline_rows: list[dict[str, object]] = []
for display, role, run in baseline_specs:
    result_path = RESULTS / "A_baselines" / "old_data_runs" / run / "results.csv"
    metrics = best_yolo_row(result_path)
    baseline_rows.append(
        {
            "model": display,
            "role": role,
            "run": run,
            "status": "已有历史结果（旧协议、val）",
            "epoch": metrics["epoch"],
            "val_mask_p": metrics["mask_p"],
            "val_mask_r": metrics["mask_r"],
            "val_mask_map50": metrics["mask_map50"],
            "val_mask_map": metrics["mask_map"],
            "val_box_map50": metrics["box_map50"],
            "val_box_map": metrics["box_map"],
            "test_mask_p": None,
            "test_mask_r": None,
            "test_mask_map50": None,
            "test_mask_map": None,
            "params": legacy_value(run, "参数量(M)"),
            "gflops": legacy_value(run, "计算量(GFLOPs)"),
            "semantic_dice": None,
            "semantic_iou": None,
            "source": metrics["results"],
            "note": "峰值行按 Mask mAP50-95 选择；未见统一 test 结果。",
        }
    )

special_baselines = [
    (
        "Mask R-CNN R50-FPN (Torchvision)",
        "经典两阶段实例分割",
        "001_6_maskrcnn_r50_fpn_seed42",
        RESULTS / "A_baselines/old_data_runs/002_retrain/maskrcnn/001_6_maskrcnn_r50_fpn_seed42",
        RESULTS / "A_baselines/old_data_runs/002_retrain/evaluation/001_6_maskrcnn_r50_fpn_seed42_test/metrics.json",
    ),
    (
        "U-Net R18 + Watershed",
        "语义转实例辅助基线",
        "E_unet_r18_watershed_seed42",
        RESULTS / "A_baselines/old_data_runs/002_retrain/unet_watershed/E_unet_r18_watershed_seed42",
        RESULTS / "A_baselines/old_data_runs/002_retrain/evaluation/E_unet_r18_watershed_seed42_test/metrics.json",
    ),
]
for display, role, run, run_dir, test_path in special_baselines:
    val_path = run_dir / "validation/best_metrics.json"
    val = json.loads(val_path.read_text(encoding="utf-8"))
    test = json.loads(test_path.read_text(encoding="utf-8"))
    baseline_rows.append(
        {
            "model": display,
            "role": role,
            "run": run,
            "status": "已有历史结果（旧协议、val+test）",
            "epoch": find_best_epoch(run_dir / "history.json"),
            "val_mask_p": val.get("mask_precision"),
            "val_mask_r": val.get("mask_recall"),
            "val_mask_map50": val.get("mask_ap_50"),
            "val_mask_map": val.get("mask_ap_50_95"),
            "val_box_map50": val.get("box_ap_50"),
            "val_box_map": val.get("box_ap_50_95"),
            "test_mask_p": test.get("mask_precision"),
            "test_mask_r": test.get("mask_recall"),
            "test_mask_map50": test.get("mask_ap_50"),
            "test_mask_map": test.get("mask_ap_50_95"),
            "params": test.get("params_m"),
            "gflops": None,
            "semantic_dice": test.get("semantic_dice"),
            "semantic_iou": test.get("semantic_iou"),
            "source": str(test_path.resolve()),
            "note": "COCO evaluator；U-Net 经 marker-controlled watershed 转为实例后评估。"
            if display.startswith("U-Net")
            else "COCO evaluator；测试集为 94 图。",
        }
    )


non_yolo_rows: list[dict[str, object]] = []
for display, role, run, run_dir, test_path in special_baselines:
    val_path = run_dir / "validation/best_metrics.json"
    metadata_path = run_dir / "run_metadata.json"
    history_path = run_dir / "history.json"
    val = json.loads(val_path.read_text(encoding="utf-8"))
    test = json.loads(test_path.read_text(encoding="utf-8"))
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    config = metadata["config"]
    framework = metadata.get("framework", {})
    prediction_path = test_path.with_name("predictions.coco.json")
    non_yolo_rows.append(
        {
            "model": display,
            "family": "Torchvision" if display.startswith("Mask R-CNN") else "SMP + Watershed",
            "role": role,
            "status": "已完成：300 epoch、val、test",
            "run": run,
            "epochs": config.get("epochs"),
            "best_epoch": find_best_epoch(history_path),
            "dataset": "citrus_prepared（旧划分）train/val/test=659/188/94",
            "framework": f'Torch {framework.get("torch", "")} / CUDA {framework.get("cuda_runtime", "")}',
            "val": val,
            "test": test,
            "params": test.get("params_m") or config.get("parameters_m"),
            "best_weight": str((run_dir / "model_best.pth").resolve()),
            "last_weight": str((run_dir / "model_last.pth").resolve()),
            "history": str(history_path.resolve()),
            "metadata": str(metadata_path.resolve()),
            "val_metrics": str(val_path.resolve()),
            "test_metrics": str(test_path.resolve()),
            "test_predictions": str(prediction_path.resolve()),
            "launcher": str(
                (
                    SERVER_DIR / "code/baseline_choice/run_maskrcnn.py"
                    if display.startswith("Mask R-CNN")
                    else SERVER_DIR / "code/baseline_choice/run_unet.py"
                ).resolve()
            ),
            "note": (
                "Mask R-CNN 使用 COCO 初始化、R50-FPN、COCO evaluator。"
                if display.startswith("Mask R-CNN")
                else "U-Net 先输出语义前景，再以 marker-controlled watershed 拆分实例；阈值0.5、min_distance=8、min_area=20。"
            ),
        }
    )

registry_path = SERVER_DIR / "code/baseline_choice/configs/baselines.yaml"
ledger_path = WORKSPACE / "4_baseline_choice/runs/_protocol/ledger.jsonl"
missing_non_yolo = [
    ("RTMDet-Ins-tiny", "MMDetection", "轻量一阶段实例分割", SERVER_DIR / "code/baseline_choice/run_mmdet.py", "dry-run 队列中出现；未发现训练目录、权重或 metrics"),
    ("SOLOv2-Light R18-FPN", "MMDetection", "无框位置式实例分割", SERVER_DIR / "code/baseline_choice/run_mmdet.py", "dry-run 队列中出现；未发现训练目录、权重或 metrics"),
    ("Mask R-CNN R50-FPN (MMDetection)", "MMDetection", "两阶段交叉验证", SERVER_DIR / "code/baseline_choice/run_mmdet.py", "只有运行入口；未发现 MMDetection 版训练结果"),
    ("RF-DETR Seg Preview @312", "RF-DETR", "Transformer 实例分割", SERVER_DIR / "code/baseline_choice/run_rfdetr.py", "dry-run 队列中出现；未发现训练目录、权重或 metrics"),
    ("DeepLabV3+ + Watershed", "SMP + Watershed", "可选语义转实例", registry_path, "只有注册项；未发现训练输出"),
    ("SegFormer-B0 + Watershed", "SMP + Watershed", "可选语义转实例", registry_path, "只有注册项；未发现训练输出"),
]
for display, family, role, launcher, note in missing_non_yolo:
    non_yolo_rows.append(
        {
            "model": display,
            "family": family,
            "role": role,
            "status": "未找到结果：仅有配置/运行入口",
            "run": "",
            "epochs": None,
            "best_epoch": None,
            "dataset": "",
            "framework": "",
            "val": {},
            "test": {},
            "params": None,
            "best_weight": "",
            "last_weight": "",
            "history": "",
            "metadata": "",
            "val_metrics": "",
            "test_metrics": "",
            "test_predictions": "",
            "launcher": str(Path(launcher).resolve()),
            "note": f"{note}。注册表：{registry_path.resolve()}；运行审计：{ledger_path.resolve()}",
        }
    )


selected_specs = [
    (
        "YOLO11n-seg",
        "基线",
        "001_3_yolo11-seg_adamw",
        RESULTS / "A_baselines/old_data_runs/001_3_yolo11-seg_adamw/results.csv",
        "C3k2 + SPPF + C2PSA + Segment",
        "以 78.64% Mask mAP50 建立阶段基线。",
    ),
    (
        "F61 TGP+TDAM",
        "纹理先验探索",
        "F61_yolo11-seg-tgp-tdam_300ep",
        RESULTS / "F_series/old_data_300ep/F61_yolo11-seg-tgp-tdam_300ep/results.csv",
        "TGP 输入纹理先验 + 两级 TDAM 纹理差分放大",
        "先验证果叶同色场景中的纹理判别思路。",
    ),
    (
        "F36 复合模型（无 DySample）",
        "多模块组合",
        "F36_yolo11-seg-ours-no-dysample_300ep",
        RESULTS / "F_series/old_data_300ep/F36_yolo11-seg-ours-no-dysample_300ep/results.csv",
        "SPDConv + DFEM + SPPF-LSKA + BiFPN + LIAM",
        "组合保细节下采样、频域增强、上下文与亮度适应。",
    ),
    (
        "F45 Edge-Nano",
        "轻量化代表",
        "F45_yolo11-seg-citrusfar-edge-nano_300ep",
        RESULTS / "F_series/old_data_300ep/F45_yolo11-seg-citrusfar-edge-nano_300ep/results.csv",
        "FasterBlock + HWDown + SPPF-LSKA + BiFPN + CSFG",
        "参数降至约 1.42M，同时把 Mask mAP50 提升到约 81.1%。",
    ),
    (
        "F56 FreqSuite",
        "频域代表",
        "F56_yolo11-seg-freqsuite_300ep",
        RESULTS / "F_series/old_data_300ep/F56_yolo11-seg-freqsuite_300ep/results.csv",
        "WTConv 骨干 + Haar 下采样 + MWCA 跨频带注意力",
        "保留边缘频率信息，并按内容选择有效频带。",
    ),
    (
        "F28 DFEM+SPD",
        "细节增强代表",
        "F28_yolo11-seg-dfem-spd_300ep",
        RESULTS / "F_series/old_data_300ep/F28_yolo11-seg-dfem-spd_300ep/results.csv",
        "三处 SPDConv + 中层 DFEM",
        "减少下采样信息损失，补偿模糊、暗弱目标的频带响应。",
    ),
    (
        "F25 SPD+EMA",
        "注意力代表",
        "F25_yolo11-seg-spd-ema_300ep",
        RESULTS / "F_series/old_data_300ep/F25_yolo11-seg-spd-ema_300ep/results.csv",
        "三处 SPDConv + P3 端 EMA 多尺度注意力",
        "保留小目标像素信息，并增强弱响应目标。",
    ),
]

selected_rows = []
for display, role, run, path, change, talk in selected_specs:
    metrics = best_yolo_row(path)
    selected_rows.append(
        {
            "model": display,
            "role": role,
            "run": run,
            "change": change,
            "talk": talk,
            "params": legacy_value(run, "参数量(M)"),
            "gflops": legacy_value(run, "计算量(GFLOPs)"),
            **metrics,
        }
    )


candidate_rows = []
candidate_dirs = [
    RESULTS / "F_series/old_data_300ep",
    RESULTS / "G_series/old_data_300ep",
    RESULTS / "N_series/old_data_300ep",
    RESULTS / "SXQ_series/old_data_300ep",
]
for directory in candidate_dirs:
    for path in directory.glob("*/results.csv"):
        metrics = best_yolo_row(path)
        if 0.800 <= float(metrics["mask_map50"]) <= 0.815:
            run = path.parent.name
            old = legacy_by_run.get(run, {})
            args_path = path.parent / "args.yaml"
            args = yaml.safe_load(args_path.read_text(encoding="utf-8")) if args_path.exists() else {}
            candidate_rows.append(
                {
                    "model": old.get("模型名称", run),
                    "run": run,
                    "series": path.relative_to(RESULTS).parts[0],
                    "epoch": metrics["epoch"],
                    "mask_p": metrics["mask_p"],
                    "mask_r": metrics["mask_r"],
                    "mask_map50": metrics["mask_map50"],
                    "mask_map": metrics["mask_map"],
                    "params": old.get("参数量(M)", ""),
                    "gflops": old.get("计算量(GFLOPs)", ""),
                    "part": old.get("改进部位", ""),
                    "design": old.get("改进方向/设计意图", ""),
                    "model_yaml": args.get("model", ""),
                    "source": metrics["results"],
                }
            )
candidate_rows.sort(key=lambda row: (row["mask_map50"], row["mask_map"]), reverse=True)


baseline_matrix = [
    ("YOLOv8n-Seg", "yolo", "核心", "已有历史结果", "正式 grouped_dedup 待跑"),
    ("YOLO11n-Seg", "yolo", "主基线", "已有历史结果", "正式需 3 seed"),
    ("YOLO26n-Seg", "yolo", "核心", "已有历史结果", "正式 grouped_dedup 待跑"),
    ("YOLO12n-Seg", "yolo", "可选", "已有历史结果", "非最少必做项"),
    ("YOLO11s-Seg", "yolo", "可选精度参照", "仅有配置", "待跑"),
    ("RTMDet-Ins-tiny", "mmdetection", "核心", "仅有配置", "待 smoke/screen/formal"),
    ("Mask R-CNN R50-FPN (Torchvision)", "torchvision", "核心", "已有历史 val+test", "正式 grouped_dedup 待跑"),
    ("Mask R-CNN R50-FPN (MMDetection)", "mmdetection", "可选", "仅有配置", "待跑"),
    ("RF-DETR Seg Preview @312", "rfdetr", "核心", "仅有配置", "待跑；速度不可与 640 直接比较"),
    ("SOLOv2-Light R18-FPN", "mmdetection", "期刊增强", "仅有配置", "待跑"),
    ("U-Net R18 + Watershed", "semantic + watershed", "辅助", "已有历史 val+test", "正式 grouped_dedup 待跑"),
    ("DeepLabV3+ + Watershed", "semantic + watershed", "可选", "仅有配置", "与 SegFormer 二选一"),
    ("SegFormer-B0 + Watershed", "semantic + watershed", "可选", "仅有配置", "与 DeepLabV3+ 二选一"),
]


structure_rows = [
    (
        "YOLO11n-seg",
        "基线",
        "C3k2 主干、SPPF 多尺度池化、C2PSA 注意力、三尺度 Segment 头",
        "为各模块改进提供统一锚点。",
        "ultralytics 官方结构 / A_baselines 历史运行",
    ),
    (
        "F61 TGP+TDAM",
        "输入端 + 浅中层骨干",
        "TGP 从亮度构造 3/7/15 尺度局部对比纹理并做可靠性门控；TDAM 用 3/7/11 邻域差分放大纹理差异。",
        "针对绿果与叶片同色、主要依靠纹理差异的问题。代码中的 TDAM 是 Texture-Difference Amplification，并非旧说明中的“拓扑解耦注意力”。",
        "F61 YAML；citrus_far.py:TGP/TDAM",
    ),
    (
        "F36 复合模型（无 DySample）",
        "骨干 + SPPF + Neck",
        "SPDConv 保留降采样像素；DFEM 做频带与暗区补偿；SPPF-LSKA 选择长程上下文；BiFPN 加权融合；LIAM 做亮度不变注意力。",
        "一次验证多种先验是否能协同；模块较多，适合当组合探索，不适合作为清晰单因素归因。",
        "F36 YAML；citrus_far.py",
    ),
    (
        "F45 Edge-Nano",
        "全网络轻量化",
        "缩减通道和重复次数，以 FasterBlock 的部分卷积替代重卷积；HWDown 保留小波高低频；SPPF-LSKA、BiFPN 和 CSFG 补充上下文与 P2 细节。",
        "把模型压到约 1.42M 参数，同时保持约 81.1% Mask mAP50，是组会最容易讲清的轻量化代表。",
        "F45 YAML；citrus_far.py:FasterBlock/HWDown/SPPF_LSKA/BiFPNConcat/CSFG",
    ),
    (
        "F56 FreqSuite",
        "骨干 + 下采样 + 中层注意力",
        "C3k2_WT 使用小波卷积获得频域大感受野；HWDown 显式保存 LL/LH/HL/HH；MWCA 做两级小波跨频带注意力与高频显著门控。",
        "让网络按目标尺度和清晰度选择频带，侧重模糊小果与边缘信息。",
        "F56 YAML；citrus_far.py:C3k2_WT/HWDown/MWCA",
    ),
    (
        "F28 DFEM+SPD",
        "骨干下采样 + P3 中层",
        "三次 SPDConv 将 2×2 空间重排到通道，减少步长卷积丢失；DFEM 用可学习频带增益、暗区补偿和局部深度卷积增强特征。",
        "直接解决小目标经连续下采样后消失，以及远处暗弱果实响应不足。",
        "F28 YAML；citrus_far.py:SPDConv/DFEM",
    ),
    (
        "F25 SPD+EMA",
        "骨干下采样 + P3 Head",
        "三次 SPDConv 保留空间信息；EMA 在高分辨率 P3 融合后进行分组的空间与通道交互。",
        "在不改变分割头主体的情况下增强小目标弱响应；本次精简表中展示值最高，为 81.409%。",
        "F25 YAML；citrus_far.py:SPDConv/EMA",
    ),
]


def style_sheet(ws, widths: dict[int, float], freeze: str, autofilter: str | None = None) -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width
    if autofilter:
        ws.auto_filter.ref = autofilter
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="222222")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY)


def title(ws, text: str, end_column: int, note: str) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    ws.cell(1, 1, text)
    ws.cell(1, 1).font = Font(name="Microsoft YaHei", size=16, bold=True, color=WHITE)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=BLUE)
    ws.cell(1, 1).alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    ws.cell(2, 1, note)
    ws.cell(2, 1).font = Font(name="Microsoft YaHei", size=10, italic=True, color="666666")
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.cell(2, 1).alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 38


def header(ws, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, 1):
        cell = ws.cell(row, column, label)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
    ws.row_dimensions[row].height = 34


wb = Workbook()
ws = wb.active
ws.title = "组会精简版"
group_headers = [
    "序号",
    "模型",
    "定位",
    "核心结构改动",
    "Params (M)",
    "GFLOPs@640",
    "峰值Epoch",
    "Mask Precision",
    "Mask Recall",
    "Mask mAP50",
    "Mask mAP50-95",
    "ΔmAP50 (百分点)",
    "ΔmAP50-95 (百分点)",
    "组会建议讲法",
    "results.csv",
    "预测图",
]
title(
    ws,
    "柑橘幼果实例分割：基线 78.64% → 阶段性约 81%",
    len(group_headers),
    "指标均为 Mask；旧协议、单 seed、验证集峰值行，仅用于组会阶段进展。当前更强结果未收入此页，正式论文须统一 grouped_dedup/test/3 seed。",
)
header(ws, 4, group_headers)
for index, row in enumerate(selected_rows, 1):
    excel_row = index + 4
    values = [
        index,
        row["model"],
        row["role"],
        row["change"],
        optional_float(row["params"]),
        optional_float(row["gflops"]),
        row["epoch"],
        row["mask_p"],
        row["mask_r"],
        row["mask_map50"],
        row["mask_map"],
        None,
        None,
        row["talk"],
        "打开结果",
        "打开预测图" if row["prediction"] else "",
    ]
    for column, value in enumerate(values, 1):
        ws.cell(excel_row, column, value)
    ws.cell(excel_row, 12, f'=IFERROR((J{excel_row}-$J$5)*100,"")')
    ws.cell(excel_row, 13, f'=IFERROR((K{excel_row}-$K$5)*100,"")')
    ws.cell(excel_row, 15).hyperlink = row["results"]
    ws.cell(excel_row, 15).style = "Hyperlink"
    if row["prediction"]:
        ws.cell(excel_row, 16).hyperlink = row["prediction"]
        ws.cell(excel_row, 16).style = "Hyperlink"
    fill = LIGHT_ORANGE if index == 1 else LIGHT_GREEN
    for cell in ws[excel_row]:
        cell.fill = PatternFill("solid", fgColor=fill)
    ws.row_dimensions[excel_row].height = 58
for row in range(5, 5 + len(selected_rows)):
    for column in (8, 9, 10, 11):
        ws.cell(row, column).number_format = "0.00%"
    for column in (12, 13):
        ws.cell(row, column).number_format = '0.00" pt"'
style_sheet(
    ws,
    {1: 7, 2: 25, 3: 17, 4: 48, 5: 12, 6: 13, 7: 11, 8: 15, 9: 13, 10: 14, 11: 16, 12: 18, 13: 21, 14: 42, 15: 14, 16: 14},
    "B5",
    f"A4:P{4 + len(selected_rows)}",
)
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "组会展示模型 Mask mAP50"
chart.y_axis.title = "Mask mAP50"
chart.y_axis.scaling.min = 0.76
chart.y_axis.scaling.max = 0.83
chart.height = 8
chart.width = 16
chart.add_data(Reference(ws, min_col=10, min_row=4, max_row=4 + len(selected_rows)), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=2, min_row=5, max_row=4 + len(selected_rows)))
ws.add_chart(chart, "B14")


ws = wb.create_sheet("基线完成结果")
baseline_headers = [
    "模型",
    "角色",
    "Run",
    "状态/评估集",
    "最佳Epoch",
    "Val Mask P",
    "Val Mask R",
    "Val Mask mAP50",
    "Val Mask mAP50-95",
    "Val Box mAP50",
    "Val Box mAP50-95",
    "Test Mask P",
    "Test Mask R",
    "Test Mask mAP50",
    "Test Mask mAP50-95",
    "Params (M)",
    "GFLOPs@640",
    "Test Semantic Dice",
    "Test Semantic IoU",
    "结果源",
    "说明",
]
title(
    ws,
    "当前文件夹内已完成的基线与历史结构对照",
    len(baseline_headers),
    "YOLO 行为 Ultralytics val 峰值；Mask R-CNN/U-Net 同时列出 COCO evaluator 的 val 与 test。空白表示本地未发现对应结果。",
)
header(ws, 4, baseline_headers)
for row_index, row in enumerate(baseline_rows, 5):
    values = [
        row["model"],
        row["role"],
        row["run"],
        row["status"],
        row["epoch"],
        row["val_mask_p"],
        row["val_mask_r"],
        row["val_mask_map50"],
        row["val_mask_map"],
        row["val_box_map50"],
        row["val_box_map"],
        row["test_mask_p"],
        row["test_mask_r"],
        row["test_mask_map50"],
        row["test_mask_map"],
        optional_float(row["params"]),
        optional_float(row["gflops"]),
        row["semantic_dice"],
        row["semantic_iou"],
        "打开结果",
        row["note"],
    ]
    for column, value in enumerate(values, 1):
        ws.cell(row_index, column, value)
    ws.cell(row_index, 20).hyperlink = str(row["source"])
    ws.cell(row_index, 20).style = "Hyperlink"
    ws.row_dimensions[row_index].height = 48
    if row["model"] == "YOLO11n-seg":
        for cell in ws[row_index]:
            cell.fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
for row in range(5, 5 + len(baseline_rows)):
    for column in range(6, 16):
        ws.cell(row, column).number_format = "0.00%"
    for column in (18, 19):
        ws.cell(row, column).number_format = "0.00%"
style_sheet(
    ws,
    {1: 31, 2: 22, 3: 40, 4: 29, 5: 11, 6: 13, 7: 13, 8: 16, 9: 18, 10: 16, 11: 18, 12: 13, 13: 13, 14: 17, 15: 19, 16: 12, 17: 13, 18: 19, 19: 18, 20: 13, 21: 45},
    "A5",
    f"A4:U{4 + len(baseline_rows)}",
)


ws = wb.create_sheet("非YOLO结果汇总")
non_yolo_headers = [
    "模型",
    "框架/路线",
    "角色",
    "状态",
    "Run",
    "训练Epoch",
    "最佳Epoch",
    "数据规模",
    "训练环境",
    "Val Mask P",
    "Val Mask R",
    "Val Mask F1",
    "Val Mask mAP50",
    "Val Mask mAP50-95",
    "Val Mask AP75",
    "Test Mask P",
    "Test Mask R",
    "Test Mask F1",
    "Test Mask mAP50",
    "Test Mask mAP50-95",
    "Test Mask AP75",
    "Test AP-S",
    "Test AP-M",
    "Test AP-L",
    "Test Box mAP50",
    "Test Box mAP50-95",
    "Test Semantic Dice",
    "Test Semantic IoU",
    "Params (M)",
    "Test Latency ms/img",
    "Test Peak VRAM MB",
    "最佳权重",
    "训练历史",
    "运行元数据",
    "Val 指标",
    "Test 指标",
    "Test 预测",
    "配置/启动脚本",
    "定位说明",
]
title(
    ws,
    "非 YOLO 基线结果与文件位置",
    len(non_yolo_headers),
    "完成结果集中在 results/A_baselines/old_data_runs/002_retrain。绿色行为已找到训练权重和指标；灰色行只找到配置或 dry-run 记录，未填造结果。",
)
header(ws, 4, non_yolo_headers)
for row_index, row in enumerate(non_yolo_rows, 5):
    val = row["val"]
    test = row["test"]
    values = [
        row["model"],
        row["family"],
        row["role"],
        row["status"],
        row["run"],
        row["epochs"],
        row["best_epoch"],
        row["dataset"],
        row["framework"],
        val.get("mask_precision"),
        val.get("mask_recall"),
        val.get("mask_f1"),
        val.get("mask_ap_50"),
        val.get("mask_ap_50_95"),
        val.get("mask_ap_75"),
        test.get("mask_precision"),
        test.get("mask_recall"),
        test.get("mask_f1"),
        test.get("mask_ap_50"),
        test.get("mask_ap_50_95"),
        test.get("mask_ap_75"),
        test.get("mask_ap_small"),
        test.get("mask_ap_medium"),
        test.get("mask_ap_large"),
        test.get("box_ap_50"),
        test.get("box_ap_50_95"),
        test.get("semantic_dice"),
        test.get("semantic_iou"),
        row["params"],
        test.get("model_latency_ms_per_image"),
        test.get("peak_vram_mb"),
        "打开权重" if row["best_weight"] else "",
        "打开历史" if row["history"] else "",
        "打开元数据" if row["metadata"] else "",
        "打开 Val 指标" if row["val_metrics"] else "",
        "打开 Test 指标" if row["test_metrics"] else "",
        "打开预测" if row["test_predictions"] else "",
        "打开配置/脚本",
        row["note"],
    ]
    for column, value in enumerate(values, 1):
        ws.cell(row_index, column, value)
    links = {
        32: row["best_weight"],
        33: row["history"],
        34: row["metadata"],
        35: row["val_metrics"],
        36: row["test_metrics"],
        37: row["test_predictions"],
        38: row["launcher"],
    }
    for column, target in links.items():
        if target:
            ws.cell(row_index, column).hyperlink = str(target)
            ws.cell(row_index, column).style = "Hyperlink"
    color = LIGHT_GREEN if str(row["status"]).startswith("已完成") else LIGHT_GRAY
    for cell in ws[row_index]:
        cell.fill = PatternFill("solid", fgColor=color)
    ws.row_dimensions[row_index].height = 68
for row in range(5, 5 + len(non_yolo_rows)):
    for column in list(range(10, 29)):
        ws.cell(row, column).number_format = "0.00%"
    for column in (29, 30, 31):
        ws.cell(row, column).number_format = "0.00"
style_sheet(
    ws,
    {
        1: 38,
        2: 23,
        3: 25,
        4: 34,
        5: 39,
        6: 12,
        7: 12,
        8: 43,
        9: 28,
        **{column: 15 for column in range(10, 29)},
        29: 13,
        30: 20,
        31: 20,
        **{column: 16 for column in range(32, 39)},
        39: 70,
    },
    "A5",
    f"A4:AM{4 + len(non_yolo_rows)}",
)


ws = wb.create_sheet("全部约81候选")
candidate_headers = [
    "模型",
    "Run",
    "系列",
    "最佳Epoch",
    "Mask P",
    "Mask R",
    "Mask mAP50",
    "Mask mAP50-95",
    "相对78.635基线 (百分点)",
    "Params (M)",
    "GFLOPs@640",
    "改进部位",
    "历史汇总中的设计说明",
    "训练时模型YAML",
    "结果源",
]
title(
    ws,
    "旧协议中 Mask mAP50 位于 80.0%–81.5% 的候选结果",
    len(candidate_headers),
    "从 F/G/N/SXQ 的 old_data_300ep results.csv 自动筛选。这里保留所有候选供内部选择，组会精简版仅选可讲清的代表。",
)
header(ws, 4, candidate_headers)
for row_index, row in enumerate(candidate_rows, 5):
    values = [
        row["model"],
        row["run"],
        row["series"],
        row["epoch"],
        row["mask_p"],
        row["mask_r"],
        row["mask_map50"],
        row["mask_map"],
        None,
        optional_float(row["params"]),
        optional_float(row["gflops"]),
        row["part"],
        row["design"],
        row["model_yaml"],
        "打开结果",
    ]
    for column, value in enumerate(values, 1):
        ws.cell(row_index, column, value)
    ws.cell(row_index, 9, f'=IFERROR((G{row_index}-0.78635)*100,"")')
    ws.cell(row_index, 15).hyperlink = row["source"]
    ws.cell(row_index, 15).style = "Hyperlink"
    ws.row_dimensions[row_index].height = 44
for row in range(5, 5 + len(candidate_rows)):
    for column in range(5, 9):
        ws.cell(row, column).number_format = "0.00%"
    ws.cell(row, 9).number_format = '0.00" pt"'
if candidate_rows:
    ws.conditional_formatting.add(
        f"G5:G{4 + len(candidate_rows)}",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
style_sheet(
    ws,
    {1: 29, 2: 58, 3: 14, 4: 12, 5: 12, 6: 12, 7: 16, 8: 18, 9: 24, 10: 12, 11: 13, 12: 22, 13: 46, 14: 54, 15: 13},
    "A5",
    f"A4:O{4 + len(candidate_rows)}",
)


ws = wb.create_sheet("基线矩阵与状态")
matrix_headers = ["模型", "范式", "论文角色", "当前本地状态", "下一步/限制"]
title(
    ws,
    "论文计划中的全部基线矩阵与本地完成状态",
    len(matrix_headers),
    "状态依据 baseline_choice/configs/baselines.yaml、对比实验方案和 results/A_baselines。已有历史结果不等于正式 grouped_dedup 结果。",
)
header(ws, 4, matrix_headers)
for row_index, row in enumerate(baseline_matrix, 5):
    for column, value in enumerate(row, 1):
        ws.cell(row_index, column, value)
    color = LIGHT_GREEN if "已有" in row[3] else LIGHT_GRAY
    for cell in ws[row_index]:
        cell.fill = PatternFill("solid", fgColor=color)
    ws.row_dimensions[row_index].height = 36
style_sheet(ws, {1: 37, 2: 27, 3: 22, 4: 28, 5: 45}, "A5", f"A4:E{4 + len(baseline_matrix)}")


ws = wb.create_sheet("模型改进说明")
structure_headers = ["模型", "改进位置", "实际代码做了什么", "设计目的与汇报判断", "核对来源"]
title(
    ws,
    "组会候选模型结构说明（按 YAML 与代码实现核对）",
    len(structure_headers),
    "说明以实际 YAML 和 citrus_far.py 为准；已修正旧汇总中 TDAM 名称与代码含义不一致的问题。",
)
header(ws, 4, structure_headers)
for row_index, row in enumerate(structure_rows, 5):
    for column, value in enumerate(row, 1):
        ws.cell(row_index, column, value)
    ws.row_dimensions[row_index].height = 82
style_sheet(ws, {1: 31, 2: 25, 3: 74, 4: 64, 5: 58}, "A5", f"A4:E{4 + len(structure_rows)}")


ws = wb.create_sheet("口径与答辩")
notes = [
    ("78 基线是什么", "YOLO11n-seg 历史运行 001_3_yolo11-seg_adamw：Mask mAP50=78.635%，对应峰值 Epoch 252；Mask mAP50-95=62.076%。"),
    ("本次建议展示上限", "F25：Mask mAP50=81.409%，较 78.635% 提升 2.774 个百分点；不要说成当前最终最好结果，应称“本次汇报展示的阶段性代表结果”。"),
    ("更稳妥的主讲模型", "F45 Edge-Nano：Mask mAP50=81.136%，参数 1.419M，较基线 2.843M 约减少一半；轻量化故事比多模块堆叠更容易讲清。"),
    ("指标口径", "表内 78/81 均指 Mask mAP50；不能与 Box mAP50 或 Mask mAP50-95 混写。峰值行按 Mask mAP50-95 选择，沿用现有汇总口径。"),
    ("协议限制", "精简表来自旧数据/旧协议、单 seed、val 峰值。orange_yolo 历史划分存在跨 split 组泄漏；这些数字只能做组会过程汇报，不能作为论文最终公平结论。"),
    ("正式论文要求", "统一 orange_yolo_grouped_dedup_20260820，主基线与最终方法 3 seed，统一 test，报告均值±标准差；跨范式基线完成后再进入论文主表。"),
    ("导师问“是不是最好”", "建议回答：目前还有更强方案在统一协议下复核，这次先汇报结构清晰、可追溯的阶段性代表结果。"),
    ("导师问“为什么有些基线空白”", "建议回答：RTMDet、RF-DETR、SOLOv2 等跨框架基线已有运行入口和配置，但本地 results 尚无正式结果，因此没有填造数值。"),
    ("Mask R-CNN/U-Net 注意", "它们已有 COCO evaluator 的 val/test 结果；U-Net 必须写成 U-Net + Watershed，不能把纯语义输出直接称为实例分割。"),
    ("结果保护", "当前更强实验没有删除，也没有改名；只从“组会精简版”中省略。完整历史仍在 1_SEVER/results。"),
]
title(ws, "指标口径、边界与现场问答", 2, "这一页用于避免组会时把指标、数据协议和“当前最好”说错。")
header(ws, 4, ["问题", "建议表述"])
for row_index, row in enumerate(notes, 5):
    ws.cell(row_index, 1, row[0])
    ws.cell(row_index, 2, row[1])
    ws.row_dimensions[row_index].height = 65
style_sheet(ws, {1: 30, 2: 112}, "A5", f"A4:B{4 + len(notes)}")


wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUT_XLSX)


base = selected_rows[0]
md_lines = [
    "# 组会汇报速读：基线 78 到阶段性 81",
    "",
    "> 本页只展示旧协议中可追溯的阶段性代表结果。当前更强模型没有列入本次组会表；正式论文仍需 grouped_dedup、统一 test 和 3 seed。",
    "",
    "## 一、建议直接放在组会 PPT 的表",
    "",
    "| 模型 | 定位 | Params (M) | GFLOPs | Mask P | Mask R | Mask mAP50 | Mask mAP50-95 | 比基线提升 |",
    "|---|---|---:|---:|---:|---:|---:|---:|---:|",
]
for row in selected_rows:
    delta = (float(row["mask_map50"]) - float(base["mask_map50"])) * 100
    md_lines.append(
        f'| {row["model"]} | {row["role"]} | {row["params"] or "—"} | {row["gflops"] or "—"} | '
        f'{float(row["mask_p"]) * 100:.2f}% | {float(row["mask_r"]) * 100:.2f}% | '
        f'**{float(row["mask_map50"]) * 100:.3f}%** | {float(row["mask_map"]) * 100:.3f}% | {delta:+.3f} pt |'
    )
md_lines.extend(
    [
        "",
        "建议主讲时保留 4 行：YOLO11n 基线、F61 纹理探索、F45 轻量化代表、F25 注意力代表。若 PPT 空间允许，再加 F56 与 F28。",
        "",
        "## 二、一分钟结果表述",
        "",
        "早期 YOLO11n-seg 基线的 Mask mAP50 为 78.635%，Mask mAP50-95 为 62.076%。围绕小目标细节丢失、果叶同色和远处目标响应弱三个问题，我们依次尝试了纹理先验、信息保持型下采样、频域增强和轻量特征融合。本次选择展示的阶段性结果达到约 81%，其中 F25 为 81.409%，比基线提高 2.774 个百分点；F45 在 1.419M 参数下达到 81.136%，更能体现轻量化价值。",
        "",
        "如果导师问是否为当前最佳，回答：**“目前还有更强方案在统一协议下复核，这次先汇报结构清晰、可追溯的阶段性代表结果。”**",
        "",
        "## 三、每个网络到底做了什么",
        "",
    ]
)
for model, location, actual, intent, _ in structure_rows:
    md_lines.extend([f"### {model}", "", f"- 改进位置：{location}", f"- 实际改动：{actual}", f"- 目的与判断：{intent}", ""])
md_lines.extend(
    [
        "## 四、全部基线完成状态",
        "",
        "| 模型 | 范式 | 角色 | 当前状态 |",
        "|---|---|---|---|",
    ]
)
for model, family, role, status, _ in baseline_matrix:
    md_lines.append(f"| {model} | {family} | {role} | {status} |")
md_lines.extend(
    [
        "",
        "当前有可读历史指标的核心基线包括 YOLOv8n、YOLO11n、YOLO12n、YOLO26n、Mask R-CNN 和 U-Net + Watershed。RTMDet-Ins-tiny、RF-DETR、SOLOv2-Light 等已有配置和运行入口，但本地 results 中没有正式数值，因此表内明确写为待跑。",
        "",
        "## 五、必须守住的口径",
        "",
        "- 78.635% 和约 81% 都是 **Mask mAP50**，不要与 Box mAP50 或 Mask mAP50-95 混用。",
        "- 这些是旧协议、单 seed、验证集阶段结果。历史 orange_yolo 划分存在跨 split 组泄漏，不能直接作为论文最终结论。",
        "- U-Net 必须写成 **U-Net + Watershed**，因为实例结果来自语义前景后的分水岭拆分。",
        "- 当前更强结果只是从本次精简展示中省略，原始结果没有改动或删除。",
        "",
        "## 六、已找到的非 YOLO 结果",
        "",
        "| 模型 | 最佳Epoch | Val Mask mAP50 / mAP50-95 | Test Mask mAP50 / mAP50-95 | Params |",
        "|---|---:|---:|---:|---:|",
        "| Mask R-CNN R50-FPN (Torchvision) | 200 | 84.319% / 73.116% | 76.506% / 64.981% | 43.976M |",
        "| U-Net R18 + Watershed | 230 | 71.245% / 57.915% | 62.593% / 49.636% | 14.328M |",
        "",
        "两组完整结果都位于 `1_SEVER/results/A_baselines/old_data_runs/002_retrain/`。RTMDet、SOLOv2、MMDetection Mask R-CNN、RF-DETR、DeepLabV3+ 和 SegFormer-B0 只找到配置、启动脚本或 dry-run 队列，没有发现训练权重和正式 metrics。Excel 的“非YOLO结果汇总”页已为每项加入可点击的权重、history、metadata、val/test metrics 和预测文件入口。",
        "",
        "## 七、证据入口",
        "",
        '- 基线原始结果：[results.csv](../results/A_baselines/old_data_runs/001_3_yolo11-seg_adamw/results.csv)',
        '- 约 81 候选目录：[old_data_300ep](../results/F_series/old_data_300ep)',
        '- 基线配置矩阵：[baselines.yaml](../code/baseline_choice/configs/baselines.yaml)',
        '- 模块实现：[citrus_far.py](../code/ultralytics-main-new/ultralytics/nn/modules/citrus_far.py)',
    ]
)
OUT_MD.write_text("\n".join(md_lines) + "\n", encoding="utf-8")


check = load_workbook(OUT_XLSX, data_only=False)
formula_errors = []
for sheet in check.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and any(error in cell.value for error in ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?")):
                formula_errors.append(f"{sheet.title}!{cell.coordinate}:{cell.value}")
if formula_errors:
    raise RuntimeError("Formula errors: " + "; ".join(formula_errors))

print(
    json.dumps(
        {
            "xlsx": str(OUT_XLSX),
            "markdown": str(OUT_MD),
            "baseline_rows": len(baseline_rows),
            "non_yolo_rows": len(non_yolo_rows),
            "near_81_candidates": len(candidate_rows),
            "formula_errors": 0,
        },
        ensure_ascii=False,
        indent=2,
    )
)
